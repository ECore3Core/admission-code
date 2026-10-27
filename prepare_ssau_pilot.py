"""Local-only pilot: general applications, public 2025 capacities, strict admission-list order."""
from pathlib import Path
from collections import defaultdict, Counter
from datetime import datetime, date
import re, json, warnings, hashlib, argparse
import pdfplumber
from openpyxl import load_workbook

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--input',type=Path,required=True)
parser.add_argument('--plan',type=Path,required=True,help='Official 2025 plan PDF, resolution 10')
parser.add_argument('--out',type=Path,required=True)
args=parser.parse_args()
OUT=args.out
DATA=args.input
def clean(x): return re.sub(r'[^а-яa-z0-9]+','',str(x or '').lower().replace('ё','е'))
def text(x): return '' if x is None else str(x).strip()
def asdate(x):
    if isinstance(x,(date,datetime)): return x.strftime('%Y-%m-%d')
    for f in ('%d.%m.%Y','%d.%m.%Y %H:%M:%S','%d.%m.%Y %H:%M'):
        try: return datetime.strptime(text(x),f).strftime('%Y-%m-%d')
        except ValueError: pass
    return None

tables=[]
with pdfplumber.open(args.plan) as pdf:
    for i,page in enumerate(pdf.pages):
        for table in page.extract_tables():
            tables.append({'page':i+1,'rows':table})
plan=[]
for tab in tables[:8]:
    for row in tab['rows']:
        if len(row)<9: continue
        code=text(row[0]); form=text(row[3]).lower()
        if re.fullmatch(r'\d\d\.\d\d\.\d\d',code) and form and text(row[4]).isdigit():
            total,separate,special,target=map(lambda x:int(x or 0),row[4:8])
            assert min(total,separate,special,target)>=0
            assert separate+special+target<=total
            plan.append({'id':f'plan-{len(plan)+1:03}', 'code':code,'form':form,
                         'names':[text(row[2])], 'total':total,
                         'capacity':total-separate-special-target,
                         'separate':separate,'special':special,'target':target,'page':tab['page']})
        elif plan and row[2] and not row[4] and not row[0]:
            plan[-1]['names'].append(text(row[2]))

warnings.filterwarnings('ignore',message='Workbook contains no default style')
wb=load_workbook(DATA,read_only=True,data_only=True)
ws=wb['Абитуриенты']; ws.reset_dimensions(); it=ws.iter_rows(values_only=True)
for row in it:
    if 'Личный номер' in row:
        headers={x:i for i,x in enumerate(row)}; break
records=defaultdict(list); stats=Counter(); unmatched=set()
for row in it:
    def get(k): return row[headers[k]] if headers[k]<len(row) else None
    if not isinstance(get('№'),(int,float)): continue
    stats['source_rows']+=1
    if get('Вид возмещения затрат')!='бюджет' or get('Вид приема')!='Общий конкурс': continue
    stats['general_budget_rows']+=1
    sid=text(get('Личный номер')); op=text(get('Набор ОП')); form=text(get('Форма обучения')).lower()
    code=op[:8]
    candidates=[p for p in plan if p['code']==code and p['form']==form]
    if len(candidates)>1:
        profile=op[op.find('(')+1:op.rfind(')')] if '(' in op and ')' in op else ''
        candidates=[p for p in candidates if clean(profile) and any(clean(profile)==clean(n) for n in p['names'])]
    if len(candidates)!=1:
        stats['unmapped_rows']+=1; unmatched.add((op,form)); continue
    p=candidates[0]
    if p['capacity']==0:
        stats['zero_capacity_rows']+=1; continue
    score=get('Сумма баллов'); priority=get('Приоритет'); d=asdate(get('Дата подачи'))
    # p.7.5 official list of applications: current sum, then unique applicant code.
    raw_code=text(get('Код в СП')); raw_code=raw_code[:-2] if raw_code.endswith('.0') else raw_code
    if not sid or not d or not isinstance(score,(int,float)) or not isinstance(priority,(int,float)) or not raw_code.isdigit():
        stats['incomplete_rows']+=1; continue
    records[sid].append({'pool':p['id'], 'priority':int(priority), 'date':d,
                         'score':int(round(score*100)), 'code':int(raw_code),
                         'original':text(get('Сдан оригинал')).lower()=='да',
                         'consent':text(get('СЗ в ПК')).lower()=='да'})
wb.close()

accepted={}; excluded=Counter()
for sid,rows in records.items():
    bypool=defaultdict(list)
    for r in rows: bypool[r['pool']].append(r)
    if any(len({r['priority'] for r in v})>1 for v in bypool.values()):
        excluded['conflicting_priority_students']+=1; continue
    if any(len({(r['score'],r['code'],r['date']) for r in v})>1 for v in bypool.values()):
        excluded['conflicting_data_students']+=1; continue
    unique=[v[0] for v in bypool.values()]
    if len({r['priority'] for r in unique})!=len(unique):
        excluded['unresolved_route_order_students']+=1; continue
    if len({(r['original'],r['consent']) for r in rows})>1:
        excluded['conflicting_confirmation_students']+=1; continue
    accepted[sid]=unique

# Anonymous labels are local only. No raw student IDs or codes are written.
order=sorted(accepted, key=lambda a:(min(r['date'] for r in accepted[a]),a))
anonymous={sid:f'a-{i+1:05}' for i,sid in enumerate(order)}
# Preserve the source unique-code order without storing those codes.
codes=sorted({r['code'] for rows in accepted.values() for r in rows})
code_rank={c:i for i,c in enumerate(codes)}
all_dates=sorted({r['date'] for rows in accepted.values() for r in rows})
current={}; batches=[]
for d in all_dates:
    events=[]
    for sid in order:
        rows=[r for r in accepted[sid] if r['date']<=d]
        if not rows: continue
        student={'id':anonymous[sid], 'original':rows[0]['original'], 'consent':rows[0]['consent'],
                 'offers':[{'pool':r['pool'],'priority':r['priority'],'route_order':0,
                            'key':[-r['score'],code_rank[r['code']]],'basis':'general-application-list'}
                           for r in sorted(rows,key=lambda x:x['priority'])]}
        if current.get(sid)!=student:
            events.append({'op':'replace' if sid in current else 'add','student':student})
            current[sid]=student
    if events:
        batches.append({'date':d,'events':events})
used={r['pool'] for rows in accepted.values() for r in rows}
pools=[{'id':p['id'],'program':p['id'],'capacity':p['capacity'],'category':'main'} for p in plan if p['id'] in used]
totals={p['id']:p['capacity'] for p in pools}
aggregate={'source_sha256':hashlib.sha256(DATA.read_bytes()).hexdigest(), 'scope':'general-budget; p7.5 application-list ranking; fixed final scores and confirmations',
           'source':dict(stats),'excluded_students':dict(excluded),'N':len(accepted),
           'E':sum(map(len,accepted.values())),'G':len(pools),'D':len(batches),
           'seats':sum(totals.values()),'unmapped_program_forms':len(unmatched),
           'raw_source_period':['2025-06-20','2025-07-07'],
           'public_plan_url':'https://ssau.ru/files/priem_doc/bachelor/2025/rasp_10_2025.pdf',
           'selected_plan':[p for p in plan if p['id'] in used]}
OUT.mkdir(parents=True,exist_ok=True)
(OUT/'pilot_input_PRIVATE.json').write_text(json.dumps({'metadata':aggregate,'pools':pools,'totals':totals,'confirmation':'consent','batches':batches},ensure_ascii=False),encoding='utf-8')
(OUT/'preparation_summary.json').write_text(json.dumps(aggregate,ensure_ascii=False,indent=2),encoding='utf-8')
# Program labels for manual capacity-map review; no applicant records.
(OUT/'unmatched_programs_LOCAL.json').write_text(json.dumps(sorted(unmatched),ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in aggregate.items() if k!='selected_plan'},ensure_ascii=False,indent=2))
