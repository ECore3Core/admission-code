"""Local-only workbook audit; output contains aggregates, never applicant records."""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, date
import hashlib
import json
from pathlib import Path
import warnings
from openpyxl import load_workbook

GROUP_COLUMNS = ('Филиал', 'Набор ОП', 'Вид возмещения затрат', 'Форма обучения',
                 'Вид приема', 'Дополнительный набор', 'Целевой прием',
                 'Отдельная квота', 'Особая квота', 'Предложение ЦП')

def norm(value):
    return '' if value is None else str(value).strip()

def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = norm(value)
    for fmt in ('%d.%m.%Y', '%d.%m.%Y %H:%M:%S', '%d.%m.%Y %H:%M', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None

def audit(path):
    warnings.filterwarnings('ignore', message='Workbook contains no default style')
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['Абитуриенты']
    ws.reset_dimensions()  # These exports incorrectly declare A1:A1 dimensions.
    iterator = ws.iter_rows(values_only=True)
    for row in iterator:
        if 'Личный номер' in row and 'Дата подачи' in row:
            headers = {norm(v): i for i, v in enumerate(row)}
            break
    else:
        raise ValueError('Missing recognized header row')
    ids, snils, applications, groups, programs, edges, dates = set(), set(), set(), set(), set(), set(), Counter()
    fields = Counter()
    cats = {k: Counter() for k in ('Приемная кампания', 'Вид возмещения затрат', 'Форма обучения', 'Вид приема')}
    key_to_snils, snils_to_ids, edge_priorities = defaultdict(set), defaultdict(set), defaultdict(set)
    by_date_ids, by_date_edges = defaultdict(set), defaultdict(set)
    budget_ids, budget_edges, budget_groups, budget_dates = set(), set(), set(), set()
    valid_ids, valid_edges, valid_groups = set(), set(), set()
    flags = {k: Counter() for k in ('Сдан оригинал','СЗ в ПК','СЗ на конкурс','Итоговое согласие')}
    rows = missing_id = bad_date = missing_priority = 0
    for row in iterator:
        def get(name):
            i = headers.get(name)
            return row[i] if i is not None and i < len(row) else None
        # Numbered rows are source records, not a title or totals row.
        if not isinstance(get('№'), (int, float)):
            continue
        rows += 1
        for name in headers:
            if norm(get(name)):
                fields[name] += 1
        for name in cats:
            cats[name][norm(get(name))] += 1
        for name in flags:
            flags[name][norm(get(name))] += 1
        sid = norm(get('Личный номер'))
        if not sid:
            missing_id += 1
            continue
        ids.add(sid)
        sn = ''.join(c for c in norm(get('СНИЛС')) if c.isdigit())
        if sn:
            snils.add(sn); key_to_snils[sid].add(sn); snils_to_ids[sn].add(sid)
        applications.add((sid, norm(get('№ заявления'))))
        group = tuple(norm(get(k)) for k in GROUP_COLUMNS)
        groups.add(group)
        programs.add(tuple(norm(get(k)) for k in ('Филиал', 'Набор ОП', 'Форма обучения', 'Вид возмещения затрат')))
        edge = (sid, group)
        edges.add(edge)
        priority = norm(get('Приоритет'))
        edge_priorities[edge].add(priority)
        missing_priority += not bool(priority)
        d = parse_date(get('Дата подачи'))
        has_route = bool(norm(get('Набор ОП')) and priority and norm(get('Вид возмещения затрат')))
        if has_route:
            valid_ids.add(sid); valid_edges.add(edge); valid_groups.add(group)
        if has_route and norm(get('Вид возмещения затрат')).lower() == 'бюджет':
            budget_ids.add(sid); budget_edges.add(edge); budget_groups.add(group)
            if d is not None:
                budget_dates.add(d.isoformat())
        if d is None:
            bad_date += 1
        else:
            dates[d.isoformat()] += 1
            by_date_ids[d.isoformat()].add(sid)
            by_date_edges[d.isoformat()].add(edge)
    result = {
        'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'created': str(wb.properties.created), 'rows': rows, 'N': len(ids),
        'unique_snils': len(snils), 'application_ids': len(applications),
        'E_observed': len(edges), 'G_observed': len(groups), 'program_forms': len(programs),
        'D_filing': len(dates), 'filing_min': min(dates) if dates else None,
        'valid_routes': {'N': len(valid_ids), 'E_observed':len(valid_edges), 'G_observed':len(valid_groups)},
        'budget': {'N':len(budget_ids), 'E_observed':len(budget_edges), 'G_observed':len(budget_groups), 'D_filing':len(budget_dates)},
        'filing_max': max(dates) if dates else None,
        'missing_id': missing_id, 'invalid_date': bad_date, 'missing_priority': missing_priority,
        'ids_with_multiple_snils': sum(len(v)>1 for v in key_to_snils.values()),
        'snils_with_multiple_ids': sum(len(v)>1 for v in snils_to_ids.values()),
        'edges_with_multiple_priorities': sum(len(v)>1 for v in edge_priorities.values()),
        'nonempty_fields': dict(fields), 'categories': {k:dict(v) for k,v in cats.items()},
        'flags': {k:dict(v) for k,v in flags.items()},
        'daily_rows': dict(sorted(dates.items())),
        'group_columns': GROUP_COLUMNS,
    }
    wb.close()
    return result, ids, edges

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    results = []
    known_ids, known_edges = set(), set()
    for path in sorted(args.directory.glob('*.xlsx'), key=lambda p: (p.name[3:5], p.name[:2])):
        result, ids, edges = audit(path)
        result['new_ids_vs_prior_files'] = len(ids-known_ids)
        result['new_edges_vs_prior_files'] = len(edges-known_edges)
        result['previous_ids_absent'] = len(known_ids-ids)
        result['previous_edges_absent'] = len(known_edges-edges)
        known_ids.update(ids); known_edges.update(edges)
        results.append(result)
        print(json.dumps({k:v for k,v in result.items() if k not in ('nonempty_fields','categories','daily_rows','group_columns')}, ensure_ascii=False), flush=True)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps({'files':results, 'union_N':len(known_ids), 'union_E_observed':len(known_edges)},ensure_ascii=False, indent=2),encoding='utf-8')

if __name__ == '__main__':
    main()
